from pathlib import Path
from docx import Document
from lxml import etree
from zipfile import ZipFile

BASE = Path(r"D:\实习工作\掌眼\周报")
FILES = [
    BASE / "模板" / "工作总结模板.docx",
    BASE / "模板" / "xxx（姓名）第x周工作周报(5)(1).docx",
    BASE / "第二周" / "李晟豪_第二周工作总结报告_掌眼项目.docx",
    BASE / "第二周" / "李晟豪_第二周工作周报_掌眼项目.docx",
]

for path in FILES:
    print(f"\n=== {path} ===")
    doc = Document(path)
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti}: rows={len(table.rows)} cols={len(table.columns)}")
        seen = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                paras = [p.text.replace("\n", "\\n") for p in cell.paragraphs]
                print(f"R{ri}C{ci}: {paras}")

with ZipFile(FILES[0]) as archive:
    root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    week_cell = root.xpath("//w:tbl[1]/w:tr[4]/w:tc[1]", namespaces=ns)[0]
    print("\nDIRECT_XML_PARAGRAPHS")
    for index, paragraph in enumerate(week_cell.xpath("./w:p", namespaces=ns)):
        print(index, "".join(paragraph.itertext()))

with ZipFile(FILES[1]) as archive:
    root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    print("\nWEEKLY_DIRECT_CELLS")
    for ri, row in enumerate(root.xpath("//w:tbl[1]/w:tr", namespaces=ns)):
        cells = row.xpath("./w:tc", namespaces=ns)
        records = []
        for ci, cell in enumerate(cells):
            span = cell.xpath("string(./w:tcPr/w:gridSpan/@w:val)", namespaces=ns) or "1"
            text = "".join(cell.itertext())
            records.append((ci, span, text))
        print(ri, records)
