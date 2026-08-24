from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from build_week3_reports import (
    SUMMARY_COMPLETION,
    SUMMARY_OUTPUT,
    SUMMARY_TASK,
    SUMMARY_TEMPLATE,
    WEEKLY_FIELDS,
    WEEKLY_OUTPUT,
    WEEKLY_TEMPLATE,
)


EXPECTED_TEMPLATE_HASHES = {
    SUMMARY_TEMPLATE: "D35B50AFB5D67F765585DBAB3E0571C0D43F1E3078D2A32BACB5240B208FDC92",
    WEEKLY_TEMPLATE: "EFEBA7817D7CE463DE81A5CC94E605488B793E76A72A7E9E07CF4F59CB6A3D61",
}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
STRUCTURE_XPATHS = [
    "//w:tblPr",
    "//w:tblGrid",
    "//w:trPr",
    "//w:tcPr",
    "//w:pPr",
    "//w:sectPr",
]


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest().upper()


def zip_parts(path: Path):
    with ZipFile(path) as archive:
        bad = archive.testzip()
        assert bad is None, f"CRC failure in {path}: {bad}"
        return {name: archive.read(name) for name in archive.namelist()}


def canonical(node) -> bytes:
    return etree.tostring(node, method="c14n")


def xml_root(parts):
    return etree.fromstring(parts["word/document.xml"])


def verify_package(template: Path, output: Path):
    template_parts = zip_parts(template)
    output_parts = zip_parts(output)
    assert set(template_parts) == set(output_parts), f"package part set changed: {output}"
    changed = []
    for name in sorted(template_parts):
        if sha256(template_parts[name]).digest() != sha256(output_parts[name]).digest():
            changed.append(name)
    assert changed == ["word/document.xml"], f"unexpected changed parts in {output}: {changed}"

    source = xml_root(template_parts)
    target = xml_root(output_parts)
    for tag in ["tbl", "tr", "tc", "p", "sectPr"]:
        assert len(source.xpath(f"//w:{tag}", namespaces=NS)) == len(
            target.xpath(f"//w:{tag}", namespaces=NS)
        ), f"w:{tag} count changed in {output}"
    for xpath in STRUCTURE_XPATHS:
        left = [canonical(node) for node in source.xpath(xpath, namespaces=NS)]
        right = [canonical(node) for node in target.xpath(xpath, namespaces=NS)]
        assert left == right, f"format structure changed at {xpath} in {output}"


def verify_summary():
    doc = Document(SUMMARY_OUTPUT)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 13 and len(table.columns) == 4
    assert table.cell(0, 1).text == "李晟豪"
    assert table.cell(0, 3).text == "《掌眼》微信端H5小游戏 / 游戏策划实习生"
    assert table.cell(1, 0).text == "项目起止时间:2026年7月14日起"

    paras = table.cell(3, 0).paragraphs
    prior = Document(SUMMARY_OUTPUT.parent.parent / "第二周" / "李晟豪_第二周工作总结报告_掌眼项目.docx")
    prior_paras = prior.tables[0].cell(3, 0).paragraphs
    assert paras[2].text == prior_paras[2].text, "first-week task changed"
    assert paras[3].text == prior_paras[3].text, "first-week completion changed"
    prior_second_heading = next(i for i, p in enumerate(prior_paras) if p.text == "第二周：")
    assert paras[7].text == prior_paras[prior_second_heading + 1].text, "second-week task changed"
    assert paras[8].text == prior_paras[prior_second_heading + 2].text, "second-week completion changed"
    assert paras[12].text == "1，布置任务：\n" + SUMMARY_TASK
    assert paras[13].text == "2，完成情况、困难和收获：\n" + SUMMARY_COMPLETION
    assert "老师提出封存当前高保真版本" in paras[13].text
    assert "更自动化、低认知门槛和低上手难度" in paras[13].text
    assert paras[17].text == "1，布置任务："
    assert paras[18].text == "2，完成情况、困难和收获："
    assert paras[21].text == "1，布置任务："
    assert paras[22].text == "2，完成情况、困难和收获："


def verify_weekly():
    doc = Document(WEEKLY_OUTPUT)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 11 and len(table.columns) == 3
    assert table.cell(1, 0).text == "姓名：李晟豪"
    assert table.cell(1, 2).text == "岗位名称：游戏策划实习生"
    assert table.cell(2, 0).text == "提交日期：2026年8月2日"
    for (row, cell), value in WEEKLY_FIELDS.items():
        visual_cell = 2 if row == 10 else cell
        assert table.cell(row, visual_cell).text == value, f"weekly field mismatch at row {row}"
    assert "老师要求封存当前版本" in table.cell(4, 2).text
    assert table.cell(7, 2).text.startswith("先封存当前高保真版本")
    assert "更自动化、低认知门槛和低上手难度" in table.cell(7, 2).text
    all_text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    assert "第三周" not in all_text or "第三周" in WEEKLY_OUTPUT.name
    assert "xxx" not in all_text and "第x周" not in all_text


def main():
    for template, expected in EXPECTED_TEMPLATE_HASHES.items():
        actual = file_hash(template)
        assert actual == expected, f"template changed: {template} {actual}"
    outputs = [SUMMARY_OUTPUT, WEEKLY_OUTPUT]
    assert all(path.exists() and path.stat().st_size > 0 for path in outputs)
    assert set(path.name for path in SUMMARY_OUTPUT.parent.glob("*.docx")) == {path.name for path in outputs}
    verify_package(SUMMARY_TEMPLATE, SUMMARY_OUTPUT)
    verify_package(WEEKLY_TEMPLATE, WEEKLY_OUTPUT)
    verify_summary()
    verify_weekly()
    print("PASS template hashes unchanged: 2/2")
    print("PASS DOCX CRC/open checks: 2/2")
    print("PASS package preservation: only word/document.xml changed")
    print("PASS template geometry/style structures preserved")
    print("PASS cumulative summary content and third-week weekly fields")


if __name__ == "__main__":
    main()
